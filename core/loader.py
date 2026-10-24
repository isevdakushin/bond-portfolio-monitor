"""Загрузка файла портфеля в «снимок» (Snapshot).

Устойчивость к другому формату файла:
- форматы: .xlsx, .xlsm, .xls (старый Excel), .csv;
- лист с позициями ищется по названию («Portfolio», «Портфель», «Позиции»...), а если такого нет —
  по содержимому: берётся лист, где лучше всего распознаются заголовки колонок;
- строка заголовков ищется по названиям колонок (русские и английские синонимы), в том числе
  двухстрочные заголовки; номера строк и колонок нигде не зашиты;
- единицы определяются по данным: проценты/доли, бп/проценты для спреда, рыночная стоимость вместо весов;
- рейтинги S&P, Fitch и Moody's, несколько колонок рейтинга -> составной рейтинг;
- строки-разделы («ST Corp Bonds IG» над группой бумаг) задают sleeve, строки «Итого» исключаются;
- всё, что не распознано автоматически, можно указать вручную (overrides): лист, колонки, дату снимка.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from . import ratings as rt

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}
_MONTH_RE = "jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec"

SHEET_KEYS = {
    "dashboard": ["dashboard", "дашборд", "summary", "сводка"],
    "portfolio": ["portfolio", "портфель", "позиции", "holdings", "positions", "состав"],
    "allocation": ["allocation", "аллокац", "распределение активов"],
    "breakdown": ["breakdown", "разбивка"],
}

# Каноническое имя колонки -> варианты заголовка (регистр, пробелы и знаки препинания не важны).
PORTFOLIO_COLUMNS = {
    "num": ["#", "№", "No", "№ п/п", "N п/п"],
    "sleeve": ["Sleeve", "Сегмент", "Asset class", "Класс активов", "Стратегия", "Strategy", "Портфельный блок"],
    "weight": ["Weight", "Вес", "Доля", "% NAV", "% of NAV", "% портфеля", "Доля в портфеле", "Weight %",
               "Вес в портфеле", "Portfolio weight"],
    "market_value": ["Market value", "MV", "Рыночная стоимость", "Стоимость", "Amount", "Сумма", "Объём",
                     "Market value USD", "Стоимость USD"],
    "isin": ["ISIN", "ISIN code", "Код ISIN"],
    "issuer": ["Issuer", "Эмитент", "Security", "Бумага", "Name", "Название", "Наименование", "Instrument",
               "Инструмент", "Description", "Bond", "Облигация", "Security name"],
    "coupon": ["Coupon", "Купон", "Cpn", "Ставка купона", "Coupon rate"],
    "maturity": ["Maturity", "Maturity date", "Погашение", "Дата погашения", "Mat date", "Срок погашения"],
    "ytm": ["YTM", "Yield to maturity", "Доходность", "Доходность к погашению", "Yield", "YTW", "Yield to worst",
            "Доходность к погашению, %"],
    "gspread": ["G-spread", "G spread", "Gspread", "G-spread, bp", "Spread to govt", "Spread to UST",
                "Спред к UST", "G-спред", "Спред", "Spread", "Govt spread"],
    "mod_duration": ["Mod. duration", "Modified duration", "Mod duration", "Дюрация", "Модифицированная дюрация",
                     "Mod dur", "MD", "Duration", "Дюрация мод."],
    "region": ["Region", "Регион"],
    "country": ["Country", "Страна", "Country of risk", "Страна риска"],
    "sector": ["Sector", "Сектор", "Industry", "Отрасль"],
    "rating": ["Rating", "Рейтинг", "S&P", "S&P rating", "Moody's", "Moodys", "Fitch", "Composite rating",
               "Кредитный рейтинг"],
    "credit_bucket": ["Credit bucket", "Bucket", "Кредитная категория"],
    "min_piece": ["Min piece", "Min lot", "Мин. лот", "Минимальный лот", "Min denomination"],
}
COLUMN_LABELS = {
    "num": "#", "sleeve": "Sleeve", "weight": "Weight", "market_value": "Market value", "isin": "ISIN",
    "issuer": "Issuer", "coupon": "Coupon", "maturity": "Maturity", "ytm": "YTM", "gspread": "G-spread",
    "mod_duration": "Mod. duration", "region": "Region", "country": "Country", "sector": "Sector",
    "rating": "Rating", "credit_bucket": "Credit bucket", "min_piece": "Min piece",
}
# Без этих полей анализ невозможен: бумага (название или ISIN), вес (или рыночная стоимость), YTM.
ESSENTIAL = [("issuer", "isin"), ("weight", "market_value"), ("ytm",)]
# Без этих полей часть анализа будет пропущена — выводится предупреждение.
IMPORTANT = ["sleeve", "maturity", "gspread", "mod_duration", "rating", "country", "sector", "region"]
RATING_EXCLUDE = ("group", "bucket", "outlook", "watch", "группа", "прогноз", "date", "дата")
TOTAL_EXACT = {"total", "итого", "итог", "всего", "subtotal", "grandtotal", "подитог", "totalportfolio"}
TOTAL_PREFIX = ("итого", "всего", "subtotal", "grandtotal", "подитог")
NONE_CHOICE = "— нет —"


class LoadError(Exception):
    """Файл нельзя разобрать (нет таблицы позиций или ключевых колонок).

    detected — что удалось распознать (листы, подписи колонок), чтобы интерфейс
    мог предложить ручное сопоставление."""

    def __init__(self, message, detected=None):
        super().__init__(message)
        self.detected = detected or {}


@dataclass
class Snapshot:
    source_name: str
    as_of: dt.date | None
    as_of_candidates: list            # [(источник, дата)]
    positions: pd.DataFrame
    allocation: pd.DataFrame | None   # sleeve, target, current, gap, portfolio_sleeve, match
    allocation_total: dict | None
    dashboard_metrics: dict           # ytm_gross/ytm_net/mod_duration/gspread -> {label,current,target,row}
    dashboard_claims: list            # [{label, text, row}]
    dashboard_texts: list             # все текстовые ячейки дашборда
    breakdown: dict                   # {"region": {name: weight}, "rating": {...}}
    file_totals: dict | None          # строка итогов листа портфеля
    sheet_names: dict
    load_notes: list = field(default_factory=list)  # [{severity, category, message}]
    detected: dict = field(default_factory=dict)    # что и где распознано (для ручной правки в интерфейсе)
    sheet_texts: dict = field(default_factory=dict)  # {лист: [текстовые ячейки]} — все листы, кроме позиций


# ---------- вспомогательные преобразования ----------

def norm(s) -> str:
    return re.sub(r"[^a-zа-яё0-9#№&]", "", str(s).lower()) if s is not None else ""


_NUM_RE = re.compile(r"^([-+−]?[\d\s .,']+)\s*(%|bp|bps|бп|б\.п\.|yrs?|years?|лет|года?)?$", re.I)


def to_float(v):
    """Число из ячейки: 5.58, '5,58%', '1 234,5', '1,234.5', '108 бп'. Проценты -> доли."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if pd.isna(v) else float(v)
    s = str(v).strip()
    if not s or s in ("-", "—", "n/a", "N/A", "#N/A", "nan"):
        return None
    m = _NUM_RE.match(s)
    if not m:
        return None
    body, unit = m.group(1), (m.group(2) or "").lower()
    body = body.replace(" ", "").replace(" ", "").replace("'", "").replace("−", "-")
    if "," in body and "." in body:     # 1,234.56 или 1.234,56: разделитель тысяч стоит первым
        body = body.replace(",", "") if body.find(",") < body.find(".") else body.replace(".", "").replace(",", ".")
    else:
        body = body.replace(",", ".")
    try:
        x = float(body)
    except ValueError:
        return None
    return x / 100 if unit == "%" else x


def to_str(v):
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    s = str(v).strip()
    return s or None


def to_date(v):
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and 20000 < v < 80000:   # серийный номер даты Excel
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v))
    s = str(v).strip()
    if not s:
        return None
    d = parse_text_date(s)
    if d:
        return d
    m = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?", s)
    if m:
        try:
            return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    ts = pd.to_datetime(s, dayfirst=True, errors="coerce")
    return None if pd.isna(ts) else ts.date()


def parse_text_date(text: str):
    """Дата из текста: '15jul2027', '14 Sep 2026', '14.09.2026', '14/09/2026'."""
    m = re.search(rf"(\d{{1,2}})\s*({_MONTH_RE})[a-z]*\.?\s*(\d{{4}})", text, re.I)
    if m:
        try:
            return dt.date(int(m.group(3)), MONTHS[m.group(2).lower()[:3]], int(m.group(1)))
        except ValueError:
            return None
    m = re.search(r"(\d{1,2})[./](\d{1,2})[./](\d{4})", text)
    if m:
        try:
            return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def parse_name_coupon(text):
    if not text:
        return None
    m = re.search(r"(\d+(?:[.,]\d+)?)\s*%", str(text))
    return float(m.group(1).replace(",", ".")) / 100 if m else None


def parse_name_maturity(text):
    if not text:
        return None
    m = re.search(rf"(\d{{1,2}})\s*({_MONTH_RE})[a-z]*\.?\s*(\d{{4}})", str(text), re.I)
    if not m:
        return None
    try:
        return dt.date(int(m.group(3)), MONTHS[m.group(2).lower()[:3]], int(m.group(1)))
    except ValueError:
        return None


# ---------- чтение файла в «сетки» значений ----------

def read_grids(data: bytes, name: str):
    """Файл -> ({имя листа: список строк}, примечания). Значения формул — сохранённые Excel'ем."""
    ext = Path(name).suffix.lower()
    notes = []
    readers = {".csv": [_read_csv], ".txt": [_read_csv], ".xls": [_read_xls, _read_xlsx]}
    for reader in readers.get(ext, [_read_xlsx, _read_xls, _read_csv]):
        try:
            return reader(data, notes), notes
        except LoadError:
            raise
        except Exception:  # noqa: BLE001 — пробуем следующий формат
            continue
    raise LoadError("Не удалось прочитать файл. Поддерживаются .xlsx, .xlsm, .xls и .csv.")


def _read_xlsx(data, notes):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), data_only=True)
    grids = {}
    for ws in wb.worksheets:
        g = [list(r) for r in ws.iter_rows(values_only=True)]
        # Объединённые ячейки (например, Sleeve на группу строк): значение — во все ячейки диапазона.
        for rng in ws.merged_cells.ranges:
            v = g[rng.min_row - 1][rng.min_col - 1] if rng.min_row - 1 < len(g) else None
            for r in range(rng.min_row - 1, min(rng.max_row, len(g))):
                for c in range(rng.min_col - 1, min(rng.max_col, len(g[r]))):
                    g[r][c] = v
        grids[ws.title] = g
    # Формулы без сохранённых значений (файл собран программой, а не сохранён в Excel) читаются пустыми.
    wbf = load_workbook(io.BytesIO(data), data_only=False)
    missing = 0
    for ws in wbf.worksheets:
        g = grids[ws.title]
        for r, row in enumerate(ws.iter_rows(values_only=True)):
            for c, v in enumerate(row):
                if isinstance(v, str) and v.startswith("=") and (r >= len(g) or c >= len(g[r]) or g[r][c] is None):
                    missing += 1
    if missing:
        notes.append({"severity": "Предупреждение", "category": "Структура",
                      "message": f"В файле {missing} формул без сохранённых значений — такие ячейки читаются пустыми. "
                                 "Откройте файл в Excel и сохраните его, чтобы значения пересчитались."})
    return grids


def _read_xls(data, notes):
    import xlrd
    book = xlrd.open_workbook(file_contents=data)
    grids = {}
    for sh in book.sheets():
        rows = []
        for r in range(sh.nrows):
            row = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                else:
                    row.append(cell.value)
            rows.append(row)
        grids[sh.name] = rows
    notes.append({"severity": "Инфо", "category": "Структура", "message": "Файл прочитан как старый Excel (.xls)."})
    return grids


def _read_csv(data, notes):
    for enc in ("utf-8-sig", "cp1251", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
        delim = dialect.delimiter
    except csv.Error:
        delim = ";" if sample.count(";") > sample.count(",") else ","
    rows = [[c if c.strip() != "" else None for c in r] for r in csv.reader(io.StringIO(text), delimiter=delim)]
    notes.append({"severity": "Инфо", "category": "Структура",
                  "message": f"Файл прочитан как CSV (разделитель «{delim}»). В CSV только один лист — "
                             "дашборд и целевая аллокация в нём отсутствуют."})
    return {"CSV": rows}


# ---------- распознавание заголовков ----------

def _match_header(raw):
    """Заголовок -> (каноническое имя, уровень совпадения 0 — точное, 1 — начало, 2 — вхождение)."""
    if raw is None or isinstance(raw, (int, float, dt.date)):
        return None, None
    s = str(raw).strip()
    n = norm(s)
    if not n:
        return None, None
    best = (None, 9)
    for key, al in PORTFOLIO_COLUMNS.items():
        for a in al:
            an = norm(a)
            if not an:
                continue
            if n == an:
                tier = 0
            elif len(an) >= 3 and n.startswith(an):
                tier = 1
            elif len(an) >= 6 and an in n:
                tier = 2
            else:
                continue
            if tier < best[1]:
                best = (key, tier)
    if best[0] == "rating" and any(x in n for x in RATING_EXCLUDE):
        return None, None
    return best if best[0] else (None, None)


def _header_labels(grid, r, two_rows):
    """Подписи колонок для строки r (или пары строк r, r+1 — двухстрочный заголовок)."""
    row = grid[r] if r < len(grid) else []
    if not two_rows:
        return [None if v is None else str(v).strip() for v in row]
    nxt = grid[r + 1] if r + 1 < len(grid) else []
    width = max(len(row), len(nxt))
    out, top = [], None
    for c in range(width):
        a = row[c] if c < len(row) else None
        b = nxt[c] if c < len(nxt) else None
        if a is not None and str(a).strip():
            top = str(a).strip()       # объединённая ячейка верхней строки действует на соседние колонки
        a_s = str(a).strip() if a is not None and str(a).strip() else None
        b_s = str(b).strip() if b is not None and str(b).strip() else None
        if b_s and _match_header(b_s)[0]:
            out.append(b_s)
        elif a_s and _match_header(a_s)[0]:
            out.append(a_s)
        elif b_s and top:
            out.append(f"{top} {b_s}")
        else:
            out.append(a_s or b_s)
    return out


def _map_labels(labels):
    colmap, ties, rating_cols = {}, {}, []
    for c, lab in enumerate(labels):
        key, tier = _match_header(lab)
        if not key:
            continue
        if key == "rating":
            rating_cols.append(c)
        if key not in colmap or tier < ties[key]:
            colmap[key], ties[key] = c, tier
    return colmap, rating_cols


def _score(colmap):
    s = len(colmap)
    ok = all(any(k in colmap for k in alt) for alt in ESSENTIAL)
    return s + (100 if ok else 0)


def _find_header(grid, max_scan=60):
    best = None
    for r in range(min(max_scan, len(grid))):
        for two in (False, True):
            labels = _header_labels(grid, r, two)
            colmap, rcols = _map_labels(labels)
            sc = _score(colmap) + (0 if not two else -0.5)   # при равенстве — однострочный заголовок
            if best is None or sc > best["score"]:
                best = {"score": sc, "row": r, "two": two, "labels": labels, "colmap": colmap, "rating_cols": rcols}
    return best


def _sheet_by_keyword(sheetnames, key):
    for name in sheetnames:
        if any(k in name.lower() for k in SHEET_KEYS[key]):
            return name
    return None


def _first_text(grid, rows=3):
    for r in range(min(rows, len(grid))):
        for v in grid[r]:
            if isinstance(v, str) and v.strip():
                return v.strip()
    return None


def _instrument_type(sleeve, bucket, issuer):
    text = " ".join(str(x) for x in (sleeve, bucket, issuer) if x).lower()
    if re.search(r"\b(fund|funds|etf|фонд)\b", text):
        return "fund"
    if re.search(r"cash|t-bill|tbill|govt|денежн|казначейск", text):
        return "cash"
    return "bond"


def _is_total(text):
    n = norm(text)
    return bool(n) and (n in TOTAL_EXACT or n.startswith(TOTAL_PREFIX))


# ---------- разбор листа портфеля ----------

def _parse_portfolio(grid, sheet, overrides, notes):
    h = _find_header(grid)
    labels, colmap, rating_cols = h["labels"], dict(h["colmap"]), list(h["rating_cols"])
    header_rows = [h["row"] + 1] + ([h["row"] + 2] if h["two"] else [])
    data_start = h["row"] + (2 if h["two"] else 1)

    # Ручные указания колонок имеют приоритет над автоматикой.
    for key, label in (overrides.get("columns") or {}).items():
        if label == NONE_CHOICE:
            colmap.pop(key, None)
            if key == "rating":
                rating_cols = []
        elif label in labels:
            colmap[key] = labels.index(label)
            if key == "rating":
                rating_cols = [labels.index(label)]

    missing = [" или ".join(COLUMN_LABELS[k] for k in alt) for alt in ESSENTIAL if not any(k in colmap for k in alt)]
    if missing:
        raise LoadError(f"На листе «{sheet}» не найдены обязательные колонки: {'; '.join(missing)}. "
                        "Укажите их вручную в блоке «Распознавание файла».",
                        detected={"sheet": sheet, "header_row": header_rows[0], "labels": [l for l in labels if l],
                                  "columns": {k: labels[c] for k, c in colmap.items()}})
    for k in IMPORTANT:
        if k not in colmap:
            notes.append({"severity": "Предупреждение", "category": "Структура",
                          "message": f"Не найдена колонка «{COLUMN_LABELS[k]}» — связанные расчёты будут пропущены "
                                     "(можно указать вручную в блоке «Распознавание файла»)."})
    if h["two"]:
        notes.append({"severity": "Инфо", "category": "Структура",
                      "message": f"Двухстрочный заголовок (строки {header_rows[0]}–{header_rows[1]})."})
    if "gspread" in colmap and norm(labels[colmap["gspread"]]) in ("spread", "спред"):
        notes.append({"severity": "Предупреждение", "category": "Структура",
                      "message": f"Тип спреда в колонке «{labels[colmap['gspread']]}» не указан — считаем, что это G-spread."})
    ytm_label = norm(labels[colmap["ytm"]])
    if "ytw" in ytm_label or "worst" in ytm_label:
        notes.append({"severity": "Предупреждение", "category": "Структура",
                      "message": f"Колонка «{labels[colmap['ytm']]}» — доходность к худшему исходу (YTW); используется как YTM."})
    if len(rating_cols) > 1:
        notes.append({"severity": "Инфо", "category": "Структура",
                      "message": "Несколько колонок рейтинга (" + ", ".join(labels[c] for c in rating_cols) +
                                 ") — используется составной рейтинг: из трёх средний, из двух худший."})

    def get(row, key):
        c = colmap.get(key)
        return row[c] if c is not None and c < len(row) else None

    rows, totals, section = [], [], None
    for r in range(data_start, len(grid)):
        row = grid[r]
        nonempty = [v for v in row if v is not None and str(v).strip() != ""]
        if not nonempty:
            continue
        issuer, isin = to_str(get(row, "issuer")), to_str(get(row, "isin"))
        w, mv = to_float(get(row, "weight")), to_float(get(row, "market_value"))
        texts = [str(v) for v in nonempty if isinstance(v, str)]
        if not isin and any(_is_total(t) for t in texts[:2]):
            if w is not None or mv is not None:
                totals.append(_totals_row(r, w, mv, row, get))
            continue
        nums = [to_float(get(row, k)) for k in ("ytm", "gspread", "mod_duration")]
        if not issuer and not isin:
            if w is not None or mv is not None:
                totals.append(_totals_row(r, w, mv, row, get))
            elif len(nonempty) == 1 and isinstance(nonempty[0], str):
                section = nonempty[0].strip()        # строка-раздел над группой бумаг
            continue
        if not isin and w is None and mv is None and all(x is None for x in nums):
            section = issuer                         # раздел, записанный в колонке названия
            continue
        if rating_cols:
            vals = [row[c] for c in rating_cols if c < len(row)]
            rating = rt.composite(vals) if len(rating_cols) > 1 else to_str(vals[0]) if vals else None
            rating_raw = " / ".join(str(v) for v in vals if v is not None)
        else:
            rating, rating_raw = None, None
        raw_mat = get(row, "maturity")
        rows.append({
            "excel_row": r + 1, "num": to_str(get(row, "num")),
            "sleeve": to_str(get(row, "sleeve")) or section,
            "weight": w, "market_value": mv, "isin": isin, "issuer": issuer or isin,
            "coupon": to_float(get(row, "coupon")),
            "maturity": to_date(raw_mat), "maturity_raw": None if raw_mat is None else str(raw_mat),
            "ytm": nums[0], "gspread": nums[1], "mod_duration": nums[2],
            "region": to_str(get(row, "region")), "country": to_str(get(row, "country")),
            "sector": to_str(get(row, "sector")),
            "rating": rt.normalize(rating) or to_str(rating), "rating_raw": rating_raw,
            "credit_bucket": to_str(get(row, "credit_bucket")), "min_piece": to_float(get(row, "min_piece")),
        })
    if not rows:
        raise LoadError(f"На листе «{sheet}» не найдено ни одной позиции под строкой заголовков.")
    df = pd.DataFrame(rows)
    if "sleeve" not in colmap and df["sleeve"].notna().any():
        notes.append({"severity": "Инфо", "category": "Структура",
                      "message": "Колонки Sleeve нет — сегменты взяты из строк-разделов над группами бумаг."})

    df, total = _normalize_units(df, totals, notes)
    notes.append({"severity": "Инфо", "category": "Структура",
                  "message": f"Лист «{sheet}», заголовки в строке {header_rows[0]}; загружено позиций: {len(df)}."})
    if total:
        notes.append({"severity": "Инфо", "category": "Структура",
                      "message": f"Строка итогов (строка {total['excel_row']}) исключена из позиций и используется "
                                 "только для сверки."})

    df["name_coupon"] = df["issuer"].map(parse_name_coupon)
    df["name_maturity"] = df["issuer"].map(parse_name_maturity)
    df["issuer_short"] = df["issuer"].map(lambda s: str(s).split(",")[0].strip() if s else s)
    df["instrument_type"] = [_instrument_type(s, b, i) for s, b, i in
                             zip(df["sleeve"], df["credit_bucket"], df["issuer"])]
    df["label"] = [f"#{n} {s}" if n else str(s) for n, s in zip(df["num"], df["issuer_short"])]
    detected = {"sheet": sheet, "header_row": header_rows[0], "labels": [l for l in labels if l],
                "columns": {k: labels[c] for k, c in colmap.items()},
                "rating_columns": [labels[c] for c in rating_cols]}
    return df, total, detected


def _totals_row(r, w, mv, row, get):
    return {"excel_row": r + 1, "weight": w, "market_value": mv, "ytm": to_float(get(row, "ytm")),
            "gspread": to_float(get(row, "gspread")), "mod_duration": to_float(get(row, "mod_duration"))}


def _normalize_units(df, totals, notes):
    # Веса: из рыночной стоимости, если колонки весов нет; проценты -> доли.
    if df["weight"].isna().all() and df["market_value"].notna().any():
        total_mv = df["market_value"].sum()
        df["weight"] = df["market_value"] / total_mv
        notes.append({"severity": "Инфо", "category": "Единицы",
                      "message": f"Колонки весов нет — веса рассчитаны из рыночной стоимости (сумма {total_mv:,.0f})."})
        for t in totals:
            t["weight"] = t["market_value"] / total_mv if t.get("market_value") else None
    wsum = df["weight"].sum(skipna=True)
    w_scale = 100 if 50 < wsum < 150 else 1
    if w_scale == 100:
        df["weight"] /= 100
        notes.append({"severity": "Инфо", "category": "Единицы", "message": "Веса заданы в процентах — приведены к долям."})
    y_scale = 1
    for col in ("ytm", "coupon"):
        med = df[col].median(skipna=True)
        if pd.notna(med) and med > 1:
            df[col] /= 100
            if col == "ytm":
                y_scale = 100
            notes.append({"severity": "Инфо", "category": "Единицы",
                          "message": f"Колонка {COLUMN_LABELS[col]} задана в процентах — приведена к долям."})
    g_scale = 1
    med = df["gspread"].median(skipna=True)
    if pd.notna(med) and med != 0:
        if abs(med) < 0.1:
            g_scale = 10000
            notes.append({"severity": "Инфо", "category": "Единицы",
                          "message": "G-spread задан в долях — приведён к базисным пунктам."})
        elif abs(med) < 15:
            g_scale = 100
            notes.append({"severity": "Инфо", "category": "Единицы",
                          "message": "G-spread задан в процентах — приведён к базисным пунктам."})
        df["gspread"] *= g_scale

    # Строка итогов: та, чей вес равен сумме весов позиций (промежуточные итоги разделов отбрасываются).
    total = None
    for t in totals:
        if t.get("weight") is not None and w_scale == 100 and t["weight"] > 1.5:
            t["weight"] /= 100
        if t.get("ytm") is not None and y_scale == 100 and t["ytm"] > 1:
            t["ytm"] /= 100
        if t.get("gspread") is not None:
            t["gspread"] *= g_scale
    candidates = [t for t in totals if t.get("weight") is not None and abs(t["weight"] - df["weight"].sum()) < 0.005]
    if candidates:
        total = candidates[-1]
    elif totals:
        total = totals[-1]
    return df, total


# ---------- дашборд, аллокация, разбивка ----------

def _parse_dashboard(grid):
    metrics, claims, texts = {}, [], []
    for row in grid:
        for v in row:
            if isinstance(v, str) and v.strip():
                texts.append(v.strip())
    texts = list(dict.fromkeys(texts))   # объединённые ячейки дают повторы
    hdr, cur_c, tgt_c = None, None, None
    for r, row in enumerate(grid):
        cur_c = tgt_c = None
        for c, v in enumerate(row):
            n = norm(v)
            if n.startswith(("текущ", "current", "факт", "actual")):
                cur_c = c
            if n.startswith(("цел", "target")):
                tgt_c = c
        if cur_c is not None and tgt_c is not None:
            hdr = r
            break
    if hdr is None:
        return metrics, claims, texts
    for r in range(hdr + 1, len(grid)):
        row = grid[r]
        label = next((str(v).strip() for v in row[:cur_c] if isinstance(v, str) and v.strip()), None)
        if not label:
            continue
        cur, tgt = _cell(grid, r, cur_c), _cell(grid, r, tgt_c)
        cur_f, tgt_f = to_float(cur), to_float(tgt)
        if cur_f is not None or (cur is None and tgt_f is not None):
            # Текущее значение может быть пустым (формула без сохранённого значения) — цель всё равно берём.
            key = _metric_key(label)
            if key and key not in metrics:
                metrics[key] = {"label": label, "current": cur_f, "target": tgt_f, "row": r + 1}
        elif isinstance(cur, str) and cur.strip():
            claims.append({"label": label, "text": cur.strip(), "row": r + 1})
    # Единицы: YTM в процентах -> доли.
    for k in ("ytm_gross", "ytm_net"):
        if k in metrics:
            for f in ("current", "target"):
                if metrics[k][f] is not None and metrics[k][f] > 1:
                    metrics[k][f] /= 100
    return metrics, claims, texts


def _metric_key(label):
    l = label.lower()
    if "fee" in l or "чист" in l or re.search(r"\bnet\b", l) or "комисс" in l:
        return "ytm_net"
    if "ytm" in l or "доходн" in l or "yield" in l:
        return "ytm_gross"
    if "durat" in l or "дюрац" in l:
        return "mod_duration"
    if "spread" in l or "спред" in l:
        return "gspread"
    return None


def _cell(grid, r, c):
    if r < len(grid) and c is not None and c < len(grid[r]):
        return grid[r][c]
    return None


SLEEVE_HEAD = ("sleeve", "сегмент", "assetclass", "классактивов", "стратегия", "strategy")


def _parse_allocation(grid):
    for r, row in enumerate(grid):
        cols = [(norm(v), c) for c, v in enumerate(row) if v is not None and not isinstance(v, (int, float))]
        sleeve_c = next((c for n, c in cols if n.startswith(SLEEVE_HEAD)), None)
        tgt_c = next((c for n, c in cols if n.startswith(("target", "цел"))), None)
        cur_c = next((c for n, c in cols if n.startswith(("current", "текущ", "факт", "actual"))), None)
        gap_c = next((c for n, c in cols if n.startswith(("gap", "отклон"))), None)
        if sleeve_c is None or tgt_c is None:
            continue
        items, total = [], None
        for rr in range(r + 1, len(grid)):
            name = to_str(_cell(grid, rr, sleeve_c))
            if not name:
                break
            rec = {"sleeve": name, "target": to_float(_cell(grid, rr, tgt_c)),
                   "current": to_float(_cell(grid, rr, cur_c)) if cur_c is not None else None,
                   "gap": to_float(_cell(grid, rr, gap_c)) if gap_c is not None else None,
                   "excel_row": rr + 1}
            if _is_total(name):
                total = rec
                break
            if rec["target"] is None:
                continue
            items.append(rec)
        df = pd.DataFrame(items)
        if df.empty:
            continue
        if df["target"].sum(skipna=True) > 1.5:
            for c in ("target", "current", "gap"):
                df[c] = df[c] / 100
        return df, total
    return None, None


def _parse_breakdown(grid):
    out = {}
    for r, row in enumerate(grid):
        for c, v in enumerate(row):
            n = norm(v)
            key = "region" if n in ("region", "регион") else "rating" if n in ("rating", "рейтинг") else None
            if not key or key in out:
                continue
            vals = {}
            for rr in range(r + 1, len(grid)):
                name = to_str(_cell(grid, rr, c))
                w = to_float(_cell(grid, rr, c + 1))
                if not name or _is_total(name) or w is None:
                    break
                vals[name] = w
            if vals and sum(vals.values()) > 1.5:
                vals = {k: x / 100 for k, x in vals.items()}
            out[key] = vals
    return out


def map_sleeves(alloc_names, port_names):
    """Сопоставление названий sleeve листа аллокации с листом портфеля.

    Сначала точное совпадение (без учёта регистра и пробелов), затем однозначное
    частичное (одно название является началом другого: 'ST Corp Bonds' -> 'ST Corp Bonds IG').
    """
    result = {}
    port_names = [p for p in dict.fromkeys(port_names) if p]
    for a in alloc_names:
        an = norm(a)
        exact = [p for p in port_names if norm(p) == an]
        if exact:
            result[a] = (exact[0], "точное")
            continue
        partial = [p for p in port_names if norm(p).startswith(an) or an.startswith(norm(p))]
        if len(partial) == 1:
            result[a] = (partial[0], "частичное")
        else:
            result[a] = (None, "неоднозначное" if partial else "не найдено")
    return result


def _filename_date(name):
    """Дата из имени файла: сначала ГГГГ-ММ-ДД, затем ДД_ММ_ГГ(ГГ)."""
    m = re.search(r"(?<!\d)(20\d{2})[_.-](\d{2})[_.-](\d{2})(?!\d)", name or "")
    if m:
        try:
            return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            pass
    m = re.search(r"(?<!\d)(\d{2})[_.-](\d{2})[_.-](\d{4}|\d{2})(?!\d)", name or "")
    if m:
        y = int(m.group(3))
        y = y + 2000 if y < 100 else y
        try:
            return dt.date(y, int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


# ---------- точка входа ----------

def load_snapshot(source, name: str | None = None, overrides: dict | None = None) -> Snapshot:
    """source — путь к файлу, байты или файловый объект (например, загрузка из Streamlit).

    overrides (необязательно): {"sheet": имя листа позиций, "columns": {канон. имя: подпись колонки},
    "as_of": date} — ручные указания, если автоматическое распознавание ошиблось.
    """
    overrides = overrides or {}
    if isinstance(source, (str, Path)):
        name = name or Path(source).name
        data = Path(source).read_bytes()
    elif isinstance(source, (bytes, bytearray)):
        data = bytes(source)
    else:
        data = source.read()
        name = name or getattr(source, "name", None)
    name = name or "portfolio.xlsx"

    grids, notes = read_grids(data, name)
    sheetnames = list(grids)

    # Лист позиций: ручной выбор -> по названию -> по содержимому (лучшее распознавание заголовков).
    port = overrides.get("sheet") if overrides.get("sheet") in grids else None
    if not port:
        kw = _sheet_by_keyword(sheetnames, "portfolio")
        scores = {s: _find_header(g)["score"] if g else -1 for s, g in grids.items()}
        if kw and scores.get(kw, -1) >= 100:
            port = kw
        else:
            # Ни один лист не распознан полностью — берём лучший; недостающие колонки
            # _parse_portfolio назовёт, и их можно будет указать вручную.
            port = kw if kw and scores[kw] == max(scores.values()) else max(scores, key=scores.get)
            if kw != port and scores[port] >= 100:
                notes.append({"severity": "Инфо", "category": "Структура",
                              "message": f"Лист с позициями определён по содержимому: «{port}»."})

    try:
        positions, totals, detected = _parse_portfolio(grids[port], port, overrides, notes)
    except LoadError as e:
        e.detected["sheets"] = sheetnames
        raise
    detected["sheets"] = sheetnames

    # Остальные листы: по названию, иначе — поиск подходящей таблицы на любом листе.
    others = [s for s in sheetnames if s != port]
    sheets = {"portfolio": port}

    def pick(key, parser, ok):
        kw = _sheet_by_keyword(others, key)
        order = ([kw] if kw else []) + [s for s in others if s != kw] + [port]
        for s in order:
            res = parser(grids[s])
            if ok(res):
                return s, res
        return None, parser([])

    sheets["dashboard"], (metrics, claims, texts) = pick("dashboard", _parse_dashboard, lambda x: bool(x[0]))
    sheets["allocation"], (allocation, alloc_total) = pick("allocation", _parse_allocation, lambda x: x[0] is not None)
    sheets["breakdown"], breakdown = pick("breakdown", _parse_breakdown, lambda x: any(x.values()))
    if sheets["dashboard"] is None:
        notes.append({"severity": "Предупреждение", "category": "Структура",
                      "message": "Не найдены целевые метрики (таблица «Текущий / Целевой») — сравнение с целью пропущено."})
    if sheets["allocation"] is None:
        notes.append({"severity": "Предупреждение", "category": "Структура",
                      "message": "Не найдена целевая аллокация (таблица Sleeve / Target) — отклонения от цели не считаются."})

    if allocation is not None and not allocation.empty:
        mapping = map_sleeves(allocation["sleeve"], positions["sleeve"])
        allocation["portfolio_sleeve"] = allocation["sleeve"].map(lambda s: mapping[s][0])
        allocation["match"] = allocation["sleeve"].map(lambda s: mapping[s][1])

    # Дата снимка: ручной ввод -> заголовки листов -> имя файла.
    candidates = []
    if overrides.get("as_of"):
        candidates.append(("указана вручную", overrides["as_of"]))
    for key in ("dashboard", "allocation", "portfolio", "breakdown"):
        s = sheets.get(key)
        if s and not any(src.endswith(f"«{s}»") for src, _ in candidates):
            t = _first_text(grids[s])
            d = parse_text_date(t) if t else None
            if d:
                candidates.append((f"заголовок листа «{s}»", d))
    fd = _filename_date(name)
    if fd:
        candidates.append(("имя файла", fd))
    as_of = candidates[0][1] if candidates else None

    sheet_texts = {s: list(dict.fromkeys(str(v).strip() for row in grids[s] for v in row
                                         if isinstance(v, str) and v.strip())) for s in others}
    return Snapshot(source_name=name, as_of=as_of, as_of_candidates=candidates, positions=positions,
                    allocation=allocation, allocation_total=alloc_total, dashboard_metrics=metrics,
                    dashboard_claims=claims, dashboard_texts=texts, breakdown=breakdown,
                    file_totals=totals, sheet_names=sheets, load_notes=notes, detected=detected,
                    sheet_texts=sheet_texts)
