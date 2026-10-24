"""Пустой шаблон файла портфеля (.xlsx) с нужными листами и колонками — без данных."""
from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .loader import COLUMN_LABELS
from .methodology import COLUMN_DOCS, COLUMN_ORDER

BOLD = Font(bold=True)
REQ_FILL = PatternFill("solid", fgColor="FFF3CD")


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "01 Portfolio"
    ws.append(["Портфель еврооблигаций, сборка ДД.ММ.ГГГГ"])
    ws.append([])
    order = ["num", "sleeve", "weight", "isin", "issuer", "coupon", "maturity", "ytm", "gspread", "mod_duration",
             "region", "country", "sector", "rating", "credit_bucket", "min_piece"]
    ws.append([COLUMN_LABELS[k] for k in order])
    for c, k in enumerate(order, start=1):
        cell = ws.cell(row=3, column=c)
        cell.font = BOLD
        if COLUMN_DOCS[k][0].startswith("Обязательно"):
            cell.fill = REQ_FILL
        ws.column_dimensions[cell.column_letter].width = max(12, len(COLUMN_LABELS[k]) + 4)

    t = wb.create_sheet("00 Dashboard")
    t.append(["Dashboard, сборка ДД.ММ.ГГГГ"])
    t.append([])
    t.append(["Key portfolio metrics", "Текущий", "Целевой"])
    for label in ("Weighted YTM (брутто)", "Чистыми (после fee)", "Weighted mod. duration", "Weighted G-spread, бп"):
        t.append([label, None, None])
    for c in range(1, 4):
        t.cell(row=3, column=c).font = BOLD
    t.column_dimensions["A"].width = 30

    a = wb.create_sheet("02 Allocation")
    a.append(["Allocation — Target vs Current, ДД.ММ.ГГГГ"])
    a.append([])
    a.append(["Sleeve", "Target", "Current"])
    for c in range(1, 4):
        a.cell(row=3, column=c).font = BOLD
    a.append(["Total", None, None])
    a.column_dimensions["A"].width = 30

    h = wb.create_sheet("Инструкция")
    h.append(["Колонка", "Обязательность", "Что содержит", "Формат / единицы"])
    for c in range(1, 5):
        h.cell(row=1, column=c).font = BOLD
    for k in COLUMN_ORDER:
        req, what, fmt, _ = COLUMN_DOCS[k]
        h.append([COLUMN_LABELS[k], req, what, fmt])
    h.append([])
    h.append(["¹ — название бумаги или ISIN; ² — вес или рыночная стоимость. Обязательные колонки выделены цветом."])
    h.append(["Листы 00 Dashboard и 02 Allocation необязательны: без них не будет сравнения с целью и аллокацией."])
    for col, w in zip("ABCD", (16, 16, 70, 45)):
        h.column_dimensions[col].width = w
    wb.move_sheet("00 Dashboard", offset=-1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
