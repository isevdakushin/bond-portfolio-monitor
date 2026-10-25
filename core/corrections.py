"""Поправки аналитика: исправления данных поверх файла компании с журналом.

- Исходный файл не меняется: поправки применяются к загруженному снимку, и все расчёты идут по исправленным данным.
- Каждая поправка хранит позицию (ISIN или название), поле, значение «было», значение «стало», основание и дату.
- Поправка применяется, только если текущее значение в файле равно «было». Если команда уже исправила файл
  (или значение изменилось), поправка не применяется и помечается — так поправки безопасно переносятся
  на следующий снимок.
- Подсказки: там, где верное значение следует из самого файла или из расчёта (дата погашения из названия,
  купон из названия, расчётная дюрация, категория по рейтингу), инструмент предлагает его сам.
"""
from __future__ import annotations

import copy
import datetime as dt
import io
import json
import re

import pandas as pd

from . import ratings
from .bondmath import modified_duration, years_between
from .loader import Snapshot, to_date, to_float

FIELDS = {   # поле -> (подпись, тип)
    "maturity": ("Погашение", "date"),
    "coupon": ("Купон", "pct"),
    "ytm": ("YTM", "pct"),
    "gspread": ("G-spread, бп", "num"),
    "mod_duration": ("Mod. duration", "num"),
    "weight": ("Вес", "pct"),
    "rating": ("Рейтинг", "rating"),
    "credit_bucket": ("Credit bucket", "text"),
    "country": ("Страна", "text"),
    "region": ("Регион", "text"),
    "sector": ("Сектор", "text"),
    "sleeve": ("Sleeve", "text"),
}


class CorrectionError(ValueError):
    pass


def pos_key(row) -> str:
    isin = row.get("isin")
    if isinstance(isin, str) and isin.strip():
        return isin.strip().upper()
    return "name:" + str(row.get("issuer") or "").strip()


def parse_value(field: str, raw):
    """Значение из ввода аналитика -> внутренний формат. Проценты: «4,5», «4.5%» или 0.045."""
    kind = FIELDS[field][1]
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        raise CorrectionError("Пустое значение.")
    if kind == "date":
        d = raw if isinstance(raw, dt.date) else to_date(raw)
        if d is None:
            raise CorrectionError(f"Не распознана дата: «{raw}». Формат: ДД.ММ.ГГГГ.")
        return d
    if kind in ("pct", "num"):
        x = raw if isinstance(raw, (int, float)) else to_float(raw)
        if x is None:
            raise CorrectionError(f"Не распознано число: «{raw}».")
        return x / 100 if kind == "pct" and abs(x) > 1 else float(x)
    if kind == "rating":
        r = ratings.normalize(raw)
        if r is None:
            raise CorrectionError(f"Рейтинг «{raw}» не распознан (S&P/Fitch или Moody's).")
        return r
    return str(raw).strip()


def fmt_value(field: str, v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return "—"
    kind = FIELDS[field][1]
    if kind == "date":
        return v.strftime("%d.%m.%Y") if isinstance(v, dt.date) else str(v)
    if kind == "pct":
        return f"{v * 100:.3f}%".replace(".", ",")
    if kind == "num":
        return f"{v:.2f}".replace(".", ",")
    return str(v)


def _to_json(field, v):
    return v.isoformat() if isinstance(v, dt.date) else v


def _from_json(field, v):
    return dt.date.fromisoformat(v) if FIELDS[field][1] == "date" and isinstance(v, str) else v


def _same(field, a, b) -> bool:
    if a is None or (isinstance(a, float) and pd.isna(a)):
        return b is None or (isinstance(b, float) and pd.isna(b))
    if FIELDS[field][1] in ("pct", "num"):
        return b is not None and not pd.isna(b) and abs(float(a) - float(b)) < 1e-9
    return a == b


def make(snap: Snapshot, label: str, field: str, new_raw, reason: str, author: str = "аналитик") -> dict:
    """Новая поправка для позиции (по подписи «#5 Targa…») с проверкой значения и основания."""
    if field not in FIELDS:
        raise CorrectionError(f"Поле «{field}» нельзя исправлять.")
    if not reason or not reason.strip():
        raise CorrectionError("Укажите основание поправки.")
    rows = snap.positions[snap.positions["label"] == label]
    if rows.empty:
        raise CorrectionError(f"Позиция «{label}» не найдена.")
    r = rows.iloc[0]
    new = parse_value(field, new_raw)
    old = r[field]
    if old is not None and not isinstance(old, str) and pd.isna(old):
        old = None
    if isinstance(old, (int, float)):
        old = float(old)
    if _same(field, old, new):
        raise CorrectionError("Новое значение совпадает с текущим.")
    return {"key": pos_key(r), "label": label, "field": field, "old": _to_json(field, old),
            "new": _to_json(field, new), "reason": reason.strip(), "author": author,
            "created": dt.date.today().isoformat()}


def apply(snap: Snapshot, corrections: list[dict]) -> tuple[Snapshot, pd.DataFrame]:
    """Применяет поправки к копии снимка. Возвращает (исправленный снимок, журнал)."""
    if not corrections:
        return snap, pd.DataFrame(columns=["Позиция", "Поле", "Было", "Стало", "Основание", "Статус"])
    out = copy.deepcopy(snap)
    p = out.positions
    keys = [pos_key(r) for _, r in p.iterrows()]
    log = []
    for c in corrections:
        field = c["field"]
        old, new = _from_json(field, c.get("old")), _from_json(field, c["new"])
        idx = [i for i, k in zip(p.index, keys) if k == c["key"]]
        rec = {"Позиция": c.get("label", c["key"]), "Поле": FIELDS[field][0], "Было": fmt_value(field, old),
               "Стало": fmt_value(field, new), "Основание": c.get("reason", ""), "Автор": c.get("author", ""),
               "Дата": c.get("created", ""), "field": field, "key": c["key"], "new_value": new}
        if not idx:
            rec["Статус"] = "не применена: позиции нет в файле"
        else:
            i = idx[0]
            cur = p.at[i, field]
            if not _same(field, cur, old):
                rec["Статус"] = f"не применена: в файле уже {fmt_value(field, cur)} (значение изменилось)"
            else:
                if p[field].dtype.kind in "fi" and not isinstance(new, (int, float)):
                    p[field] = p[field].astype(object)
                p.at[i, field] = new
                rec["Статус"] = "применена"
                rec["Позиция"] = p.at[i, "label"]
        log.append(rec)
    out.positions = p
    out.corrected_fields = {r["field"] for r in log if r["Статус"] == "применена"}
    if any(r["Статус"] == "применена" for r in log):
        out.load_notes = list(out.load_notes) + [{
            "severity": "Инфо", "category": "Поправки аналитика",
            "message": f"Применено поправок аналитика: {sum(r['Статус'] == 'применена' for r in log)}. "
                       "Исходный файл не изменён; расчёты — по исправленным данным."}]
    return out, pd.DataFrame(log)


def suggest(snap: Snapshot) -> list[dict]:
    """Подсказки: значения, которые следуют из самого файла или из расчёта. Аналитик решает, применять ли."""
    out = []
    for _, r in snap.positions.iterrows():
        lab = r["label"]
        mat, nm = r["maturity"], r.get("name_maturity")
        if r["instrument_type"] == "bond" and nm is not None and not pd.isna(nm) and mat is not None \
                and not pd.isna(mat) and nm != mat:
            reason = f"дата погашения в названии бумаги ({nm:%d.%m.%Y})"
            if snap.as_of and pd.notna(r["mod_duration"]):
                md_n = modified_duration(snap.as_of, nm, r["coupon"], r["ytm"])
                md_c = modified_duration(snap.as_of, mat, r["coupon"], r["ytm"])
                if md_n and md_c:
                    reason += (f"; расчётная дюрация по ней {md_n:.2f}, по колонке {md_c:.2f}, в файле "
                               f"{r['mod_duration']:.2f}").replace(".", ",")
                    if abs(md_c - r["mod_duration"]) < abs(md_n - r["mod_duration"]):
                        reason += " — дюрация ближе к дате из колонки, проверьте"
            out.append({"label": lab, "field": "maturity", "new": nm, "reason": reason})
        nc, cc = r.get("name_coupon"), r["coupon"]
        if nc is not None and cc is not None and not pd.isna(nc) and not pd.isna(cc) and abs(nc - cc) > 1e-4:
            out.append({"label": lab, "field": "coupon", "new": nc, "reason": "купон в названии бумаги"})
        rt, bucket = r["rating"], r.get("credit_bucket")
        if ratings.normalize(rt) and isinstance(bucket, str):
            b = ratings.parse_bucket(bucket)
            letter = ratings.letter(rt)
            if b["grades"] and letter not in b["grades"]:
                new_b = re.sub(r"\b(AAA|AA|A|BBB|BB|B|CCC)\b", letter, bucket, count=1)
                if new_b != bucket:
                    out.append({"label": lab, "field": "credit_bucket", "new": new_b,
                                "reason": f"привести категорию к рейтингу {rt} (или исправить рейтинг, если неверен он)"})
        if snap.as_of and r["instrument_type"] == "bond" and mat is not None and not pd.isna(mat) \
                and pd.notna(r["mod_duration"]) and mat > snap.as_of and r["mod_duration"] > years_between(snap.as_of, mat) + 0.01:
            md = modified_duration(snap.as_of, mat, r["coupon"], r["ytm"])
            if md:
                out.append({"label": lab, "field": "mod_duration", "new": round(md, 2),
                            "reason": "дюрация в файле больше срока до погашения; расчётная mod. duration по купону, "
                                      "YTM и погашению (буллет, без опционов)"})
    return out


def to_json(corrections: list[dict]) -> str:
    return json.dumps(corrections, ensure_ascii=False, indent=2)


def from_json(text: str) -> list[dict]:
    data = json.loads(text)
    if not isinstance(data, list) or not all(isinstance(c, dict) and {"key", "field", "new"} <= set(c) for c in data):
        raise CorrectionError("Файл поправок повреждён или не того формата.")
    for c in data:
        if c["field"] not in FIELDS:
            raise CorrectionError(f"Неизвестное поле в файле поправок: {c['field']}")
    return data


def corrected_excel(source: bytes, snap: Snapshot, log: pd.DataFrame) -> bytes | None:
    """Копия исходного .xlsx: применённые поправки записаны в ячейки (подсвечены, с примечанием),
    журнал — на отдельном листе. Для .xls и .csv — None (выгружается только журнал)."""
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    try:
        wb = load_workbook(io.BytesIO(source))
    except Exception:  # noqa: BLE001
        return None
    det = snap.detected
    ws = wb[det["sheet"]]
    colidx = det.get("colindex", {})
    fill = PatternFill("solid", fgColor="FFF3CD")
    applied = log[log["Статус"] == "применена"] if not log.empty else log
    rows = {pos_key(r): int(r["excel_row"]) for _, r in snap.positions.iterrows()}
    for _, rec in applied.iterrows():
        f, key = rec["field"], rec["key"]
        if f not in colidx or key not in rows:
            continue
        cell = ws.cell(row=rows[key], column=colidx[f] + 1)
        cell.value = rec["new_value"]
        cell.fill = fill
        cell.comment = Comment(f"Поправка аналитика: было {rec['Было']}. Основание: {rec['Основание']}", "аналитик")
    lg = wb.create_sheet("Журнал поправок")
    head = ["Позиция", "Поле", "Было", "Стало", "Основание", "Автор", "Дата", "Статус"]
    lg.append(head)
    for c in range(1, len(head) + 1):
        lg.cell(row=1, column=c).font = Font(bold=True)
    for _, rec in log.iterrows():
        lg.append([rec.get(h, "") for h in head])
    for col, w in zip("ABCDEFGH", (34, 16, 14, 14, 70, 12, 12, 30)):
        lg.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
