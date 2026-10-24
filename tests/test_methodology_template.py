"""Методология описывает все колонки; шаблон файла распознаётся загрузчиком; объединённые ячейки читаются."""
import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from core.config import load_config
from core.loader import COLUMN_LABELS, PORTFOLIO_COLUMNS, LoadError, load_snapshot
from core.methodology import COLUMN_DOCS, markdown
from core.template import build_template

FILE = Path(__file__).resolve().parent.parent / "Main_Portfolio_14_09_26_v2.xlsx"


def test_every_column_documented():
    assert set(COLUMN_DOCS) == set(PORTFOLIO_COLUMNS)
    md = markdown(load_config())
    for k in PORTFOLIO_COLUMNS:
        assert f"**{COLUMN_LABELS[k]}**" in md
    assert "Требования к загружаемому файлу" in md and "Как пользоваться" in md


def test_template_headers_recognized():
    wb = load_workbook(io.BytesIO(build_template()))
    assert {"00 Dashboard", "01 Portfolio", "02 Allocation", "Инструкция"} <= set(wb.sheetnames)
    # Пустой шаблон — нет позиций, но все колонки распознаны.
    with pytest.raises(LoadError, match="ни одной позиции"):
        load_snapshot(build_template(), name="template.xlsx")


@pytest.mark.skipif(not FILE.exists(), reason="нет файла портфеля")
def test_merged_sleeve_cells():
    """Sleeve записан один раз на группу строк в объединённой ячейке."""
    base = load_snapshot(FILE).positions
    wb = Workbook()
    ws = wb.active
    ws.append(["Sleeve", "Issuer", "Weight", "YTM", "G-spread", "Mod. duration"])
    row = 2
    for sleeve, g in base.groupby("sleeve", sort=False):
        start = row
        for i, (_, r) in enumerate(g.iterrows()):
            ws.append([sleeve if i == 0 else None, r["issuer"], r["weight"], r["ytm"], r["gspread"], r["mod_duration"]])
            row += 1
        if row - 1 > start:
            ws.merge_cells(start_row=start, start_column=1, end_row=row - 1, end_column=1)
    buf = io.BytesIO()
    wb.save(buf)
    s = load_snapshot(buf.getvalue(), name="merged.xlsx")
    assert list(s.positions["sleeve"]) == list(base["sleeve"])
