"""Устойчивость загрузчика к другому формату файла.

Из реального файла компании генерируются варианты с другими листами, заголовками, единицами,
структурой и форматом (.xlsx, .xls, .csv). На каждом варианте должны получаться те же портфельные
метрики, что и на исходном файле.
"""
import csv
import datetime as dt
import io
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from core.analytics import analyze
from core.config import load_config
from core.loader import LoadError, load_snapshot
from core.quality import run_checks
from core.ratings import MOODYS

FILE = Path(__file__).resolve().parent.parent / "Main_Portfolio_14_09_26_v2.xlsx"
pytestmark = pytest.mark.skipif(not FILE.exists(), reason="нет файла портфеля")
TO_MOODYS = {v: k for k, v in MOODYS.items()}


@pytest.fixture(scope="module")
def base():
    return load_snapshot(FILE)


def _metrics(snap):
    p = snap.positions
    return (round((p.weight * p.ytm).sum(), 9), round((p.weight * p.mod_duration).sum(), 6),
            round((p.weight * p.gspread).sum(), 4), round(p.weight.sum(), 9))


def _same(base, snap, check_sleeve=True):
    assert len(snap.positions) == 28
    assert _metrics(snap) == _metrics(base)
    assert list(snap.positions["rating"]) == list(base.positions["rating"])
    assert list(snap.positions["maturity"]) == list(base.positions["maturity"])
    if check_sleeve:
        assert list(snap.positions["sleeve"]) == list(base.positions["sleeve"])
    # аналитика и проверки отрабатывают без падений
    cfg = load_config()
    run_checks(snap, cfg)
    analyze(snap, cfg)


def _xlsx(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _d(x):
    return x.strftime("%d.%m.%Y") if isinstance(x, dt.date) else None


def test_russian_headers_percent_units_text_dates(base):
    """Один лист «Позиции», русские заголовки в другом порядке, всё в процентах, даты текстом."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Позиции"
    ws.append(["Наименование", "Доля в портфеле, %", "Доходность к погашению, %", "Спред к UST, %",
               "Дюрация", "Дата погашения", "Купон", "Рейтинг", "Сегмент", "Страна", "Отрасль", "Регион", "ISIN"])
    for _, r in base.positions.iterrows():
        ws.append([r.issuer, r.weight * 100, r.ytm * 100, r.gspread / 100, r.mod_duration, _d(r.maturity),
                   None if r.coupon != r.coupon else r.coupon * 100, r.rating, r.sleeve, r.country, r.sector,
                   r.region, r["isin"]])
    s = load_snapshot(_xlsx(wb), name="portfolio_2026-10-14.xlsx")
    _same(base, s)
    assert s.as_of == dt.date(2026, 10, 14)            # дата из имени файла в формате ГГГГ-ММ-ДД


def test_market_value_and_two_rating_agencies(base):
    """Английские синонимы, рыночная стоимость вместо весов, рейтинги S&P и Moody's в разных колонках."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Holdings"
    ws.append(["Security name", "ISIN code", "Market value USD", "Yield to maturity", "Spread to UST",
               "Modified duration", "Maturity date", "S&P", "Moody's", "Asset class", "Country of risk"])
    for _, r in base.positions.iterrows():
        ws.append([r.issuer, r["isin"], r.weight * 17_600_000, r.ytm, r.gspread, r.mod_duration, r.maturity,
                   r.rating, TO_MOODYS.get(r.rating), r.sleeve, r.country])
    s = load_snapshot(_xlsx(wb), name="holdings.xlsx")
    _same(base, s)
    assert any("составной рейтинг" in n["message"] for n in s.load_notes)


def test_sections_subtotals_two_row_header(base):
    """Нет колонки Sleeve: сегменты — строками-разделами; промежуточные итоги; двухстрочный заголовок;
    заголовок таблицы не в первой строке."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    ws.append(["Портфель еврооблигаций по состоянию на 14.09.2026"])
    ws.append([])
    ws.append(["Бумага", "Вес", "Доходность и риск", None, None, "Погашение", "Рейтинг"])
    ws.append([None, None, "YTM", "G-spread", "Mod. duration", None, None])
    for sleeve, g in base.positions.groupby("sleeve", sort=False):
        ws.append([sleeve])
        for _, r in g.iterrows():
            ws.append([r.issuer, r.weight, r.ytm, r.gspread, r.mod_duration, r.maturity, r.rating])
        ws.append(["Итого по сегменту", g.weight.sum()])
    ws.append(["Итого", base.positions.weight.sum(), (base.positions.weight * base.positions.ytm).sum()])
    s = load_snapshot(_xlsx(wb), name="report.xlsx")
    _same(base, s)
    assert s.as_of == dt.date(2026, 9, 14)
    assert abs(s.file_totals["weight"] - 1) < 1e-9          # выбран общий итог, а не промежуточный


def test_csv_semicolon_comma_decimals(base):
    """CSV из русского Excel: разделитель «;», десятичная запятая, проценты и «бп» в ячейках."""
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(["ISIN", "Эмитент", "Вес", "YTM", "G-spread", "Mod. duration", "Maturity", "Rating", "Sleeve"])
    for _, r in base.positions.iterrows():
        w.writerow([r["isin"] or "", r.issuer, f"{r.weight * 100:.4f}%".replace(".", ","),
                    f"{r.ytm * 100:.4f}%".replace(".", ","), f"{r.gspread:.2f} бп".replace(".", ","),
                    f"{r.mod_duration:.4f}".replace(".", ","), _d(r.maturity) or "", r.rating, r.sleeve])
    s = load_snapshot(out.getvalue().encode("cp1251"), name="portfolio.csv")
    _same(base, s)


def test_old_xls_format(base):
    xlwt = pytest.importorskip("xlwt")
    book = xlwt.Workbook()
    sh = book.add_sheet("Portfolio")
    heads = ["Issuer", "ISIN", "Weight", "YTM", "G-spread", "Mod. duration", "Maturity", "Rating", "Sleeve"]
    for c, h in enumerate(heads):
        sh.write(0, c, h)
    for i, (_, r) in enumerate(base.positions.iterrows(), start=1):
        vals = [r.issuer, r["isin"] or "", r.weight, r.ytm, r.gspread, r.mod_duration, _d(r.maturity) or "",
                r.rating, r.sleeve]
        for c, v in enumerate(vals):
            sh.write(i, c, v)
    buf = io.BytesIO()
    book.save(buf)
    _same(base, load_snapshot(buf.getvalue(), name="portfolio.xls"))


def test_sheets_renamed_without_keywords(base):
    """Листы переименованы в «Лист1…Лист4» — лист позиций, дашборд и аллокация находятся по содержимому."""
    wb = load_workbook(FILE)
    for i, ws in enumerate(wb.worksheets, start=1):
        ws.title = f"Лист{i}"
    s = load_snapshot(_xlsx(wb), name="file.xlsx")
    _same(base, s)
    assert s.dashboard_metrics["mod_duration"]["target"] == base.dashboard_metrics["mod_duration"]["target"]
    assert s.allocation is not None and len(s.allocation) == 6


def test_manual_column_override(base):
    """Нераспознанная колонка («Доход, %») — ошибка с подсказкой; после ручного указания — загрузка проходит."""
    wb = Workbook()
    ws = wb.active
    ws.append(["Issuer", "Weight", "Доход, %", "G-spread", "Mod. duration", "Maturity", "Rating", "Sleeve"])
    for _, r in base.positions.iterrows():
        ws.append([r.issuer, r.weight, r.ytm * 100, r.gspread, r.mod_duration, r.maturity, r.rating, r.sleeve])
    data = _xlsx(wb)
    with pytest.raises(LoadError, match="вручную"):
        load_snapshot(data, name="x.xlsx")
    s = load_snapshot(data, name="x.xlsx", overrides={"columns": {"ytm": "Доход, %"}})
    _same(base, s)


def test_manual_as_of_override(base):
    s = load_snapshot(FILE, overrides={"as_of": dt.date(2026, 10, 14)})
    assert s.as_of == dt.date(2026, 10, 14)

