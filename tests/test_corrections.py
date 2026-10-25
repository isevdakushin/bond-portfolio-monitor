"""Поправки аналитика: подсказки, применение, журнал, перенос на новый снимок, выгрузка Excel."""
import copy
import datetime as dt
import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core import corrections as C
from core.config import load_config
from core.loader import load_snapshot
from core.quality import run_checks

FILE = Path(__file__).resolve().parent.parent / "Main_Portfolio_14_09_26_v2.xlsx"
pytestmark = pytest.mark.skipif(not FILE.exists(), reason="нет файла портфеля")


@pytest.fixture(scope="module")
def snap():
    return load_snapshot(FILE)


def _fix_all(snap):
    keep = ("#5 Targa Resources Corp", "#2 AbbVie", "#13 Adani Ports & SEZ")
    return [C.make(snap, s["label"], s["field"], s["new"], s["reason"]) for s in C.suggest(snap) if s["label"] in keep]


def test_suggestions(snap):
    s = {(x["label"], x["field"]): x for x in C.suggest(snap)}
    assert s[("#5 Targa Resources Corp", "maturity")]["new"] == dt.date(2031, 4, 15)
    assert s[("#2 AbbVie", "credit_bucket")]["new"] == "A / IG"
    assert s[("#13 Adani Ports & SEZ", "mod_duration")]["new"] == pytest.approx(4.72, abs=0.01)


def test_apply_and_log_original_untouched(snap):
    fixed, log = C.apply(snap, _fix_all(snap))
    assert (log["Статус"] == "применена").all()
    t = fixed.positions.set_index("label")
    assert t.loc["#5 Targa Resources Corp", "maturity"] == dt.date(2031, 4, 15)
    assert snap.positions.set_index("label").loc["#5 Targa Resources Corp", "maturity"] == dt.date(2033, 3, 15)
    targa = log[log["Позиция"] == "#5 Targa Resources Corp"].iloc[0]
    assert (targa["Было"], targa["Стало"]) == ("15.03.2033", "15.04.2031")


def test_errors_resolved_and_totals_downgraded(snap):
    fixed, _ = C.apply(snap, _fix_all(snap))
    q = run_checks(fixed, load_config())
    assert (q["Уровень"] == "Ошибка").sum() == 0
    assert q["Комментарий"].str.contains("до поправок аналитика").any()       # итоги файла — «Инфо» с пояснением


def test_stale_correction_not_applied(snap):
    cs = _fix_all(snap)
    already_fixed = copy.deepcopy(snap)                          # «команда исправила файл к следующему месяцу»
    p = already_fixed.positions
    p.loc[p["label"] == "#5 Targa Resources Corp", "maturity"] = dt.date(2031, 4, 15)
    _, log = C.apply(already_fixed, cs)
    st = log.set_index("Позиция").loc["#5 Targa Resources Corp", "Статус"]
    assert st.startswith("не применена") and "15.04.2031" in st


def test_json_roundtrip(snap):
    cs = _fix_all(snap)
    back = C.from_json(C.to_json(cs))
    _, log = C.apply(snap, back)
    assert (log["Статус"] == "применена").all()
    with pytest.raises(C.CorrectionError):
        C.from_json('[{"key": "x"}]')


def test_validation(snap):
    with pytest.raises(C.CorrectionError, match="основание"):
        C.make(snap, "#2 AbbVie", "rating", "A-", "")
    with pytest.raises(C.CorrectionError, match="Рейтинг"):
        C.make(snap, "#2 AbbVie", "rating", "XYZ", "тест")
    c = C.make(snap, "#2 AbbVie", "ytm", "5,1", "уточнённая котировка")   # проценты вводятся как в Excel
    assert c["new"] == pytest.approx(0.051)


def test_corrected_excel(snap):
    fixed, log = C.apply(snap, _fix_all(snap))
    data = C.corrected_excel(FILE.read_bytes(), snap, log)
    wb = load_workbook(io.BytesIO(data))
    assert "Журнал поправок" in wb.sheetnames
    ws = wb["01 Portfolio"]
    targa_row = int(snap.positions.set_index("label").loc["#5 Targa Resources Corp", "excel_row"])
    mat_col = snap.detected["colindex"]["maturity"] + 1
    assert ws.cell(targa_row, mat_col).value.date() == dt.date(2031, 4, 15)
    assert ws.cell(targa_row, mat_col).comment is not None
    # исправленный файл снова читается инструментом, и ошибок по Targa нет
    q = run_checks(load_snapshot(data, name="corrected.xlsx"), load_config())
    assert not ((q["Уровень"] == "Ошибка") & q["Позиция"].str.contains("Targa")).any()


def test_report_contains_corrections_and_registry_holds(snap):
    import re

    from core.pipeline import run
    from core.report import build
    fixed, log = C.apply(snap, _fix_all(snap))
    res = run(fixed, load_config(), use_market=False, today=dt.date(2026, 10, 25))
    res.corrections = log
    rep = build(res)
    assert "Поправки аналитика" in rep.text and "15.03.2033 → 15.04.2031" in rep.text
    text = re.sub(r"\b\d{2}\.\d{2}(?:\.\d{4})?\b", " ", rep.text)
    quotes = " ".join(rep.quotes)
    bad = [t for t in re.findall(r"\d+(?:,\d+)?", text)
           if t not in rep.numbers and t not in quotes and not re.search(rf"#{t}\b", rep.text)]
    assert bad == [], f"числа вне расчёта: {bad}"


def test_pipeline_applies_corrections_and_keeps_raw(snap):
    from core.pipeline import run
    res = run(snap, load_config(), use_market=False, corrections=_fix_all(snap))
    assert res.snap_raw is snap                                           # исходный снимок — для подсказок
    assert res.snap.positions.set_index("label").loc["#5 Targa Resources Corp", "maturity"] == dt.date(2031, 4, 15)
    assert (res.corrections["Статус"] == "применена").all()
    assert (res.quality["Уровень"] == "Ошибка").sum() == 0                 # проверки — по исправленным данным
    assert run(snap, load_config(), use_market=False).corrections is None  # без поправок журнала нет
