"""СИНТЕТИЧЕСКИЙ «следующий снимок» для тестов сравнения. Не рыночные данные.

Из реального файла 14.09.2026 делается файл «14.10.2026» с заранее известными изменениями:
- IBHJ 4,5% → 8% NAV; Apple 3% → 5,5%; T-bills 16% → 10%;
- Pertamina заменена новой бумагой «SYNTH TestCo, 5% 15jun2031»;
- Adani Ports: рейтинг BBB- → BB+, G-spread +60 бп сверх общего;
- у всех облигаций YTM +30 бп, G-spread +10 бп.
"""
import io
from pathlib import Path

from openpyxl import load_workbook

FILE = Path(__file__).resolve().parent.parent / "Main_Portfolio_14_09_26_v2.xlsx"


def build_next_month() -> bytes:
    wb = load_workbook(FILE)
    wb["00 Dashboard"]["A1"] = "Dashboard — Main portfolio, сборка 14.10.2026"
    wb["02 Allocation"]["A1"] = "02 Allocation — Target vs Current (as of 14 Oct 2026)"
    wb["03 Breakdown"]["A1"] = "03 Breakdown — Region and rating distribution (% of NAV), 14.10.2026"
    ws = wb["01 Portfolio"]
    hdr = {ws.cell(6, c).value: c for c in range(1, ws.max_column + 1)}
    for r in range(7, 35):
        issuer = ws.cell(r, hdr["Issuer"]).value or ""
        sleeve = ws.cell(r, hdr["Sleeve"]).value or ""
        if "Fund" not in sleeve and "Cash" not in sleeve:
            ws.cell(r, hdr["YTM"]).value += 0.003
            ws.cell(r, hdr["G-spread"]).value += 10
        if issuer.startswith("IBHJ"):
            ws.cell(r, hdr["Weight"]).value = 0.08
        if issuer.startswith("Apple"):
            ws.cell(r, hdr["Weight"]).value = 0.055
        if issuer.startswith("US T-bills"):
            ws.cell(r, hdr["Weight"]).value = 0.10
        if issuer.startswith("Adani"):
            ws.cell(r, hdr["Rating"]).value = "BB+"
            ws.cell(r, hdr["G-spread"]).value += 60
        if issuer.startswith("Pertamina"):
            ws.cell(r, hdr["Issuer"]).value = "SYNTH TestCo, 5% 15jun2031"
            ws.cell(r, hdr["ISIN"]).value = "XS0000000000"
            ws.cell(r, hdr["Coupon"]).value = 0.05
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
